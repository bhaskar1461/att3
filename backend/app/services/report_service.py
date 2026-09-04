import io
import csv
import pandas as pd
from typing import List, Dict, Any, Optional
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

class ReportService:
    @staticmethod
    def generate_excel_report(data: List[Dict[str, Any]], title: str = "Attendance Report") -> bytes:
        """
        Generates clean formatted Excel report from attendance records.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Summary"

        # Title Block
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"COLLEGE ATTENDANCE MANAGEMENT SYSTEM — {title.upper()}"
        title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Generated Timestamp
        ws.merge_cells("A2:G2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Report Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        sub_cell.font = Font(name="Arial", size=10, italic=True, color="64748B")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Table Headers
        headers = ["S.No", "Roll Number", "Student Name", "Department", "Section", "Subject", "Status", "Date"]
        ws.append([]) # Row 3 blank
        ws.append(headers) # Row 4 headers

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        ws.row_dimensions[4].height = 25

        # Populate Rows
        for i, item in enumerate(data, 1):
            row_data = [
                i,
                item.get("roll_number", ""),
                item.get("student_name", ""),
                item.get("department", ""),
                item.get("section", ""),
                item.get("subject", ""),
                item.get("status", ""),
                item.get("date", "")
            ]
            ws.append(row_data)
            current_row = 4 + i
            ws.row_dimensions[current_row].height = 20

            # Alternate row background
            fill_color = "F8FAFC" if i % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                if col_num in [1, 2, 7, 8]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                # Highlight Present/Absent
                if col_num == 7:
                    if str(cell.value).upper() in ["PRESENT", "4"]:
                        cell.font = Font(name="Arial", size=10, bold=True, color="16A34A")
                    else:
                        cell.font = Font(name="Arial", size=10, bold=True, color="DC2626")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_csv_report(data: List[Dict[str, Any]]) -> str:
        """
        Generates CSV format attendance report.
        """
        output = io.StringIO()
        fieldnames = ["S.No", "Roll Number", "Student Name", "Department", "Section", "Subject", "Status", "Date"]
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()

        for i, item in enumerate(data, 1):
            writer.writerow({
                "S.No": i,
                "Roll Number": item.get("roll_number", ""),
                "Student Name": item.get("student_name", ""),
                "Department": item.get("department", ""),
                "Section": item.get("section", ""),
                "Subject": item.get("subject", ""),
                "Status": item.get("status", ""),
                "Date": item.get("date", "")
            })
        return output.getvalue()

    @staticmethod
    def generate_pdf_report(data: List[Dict[str, Any]], title: str = "Attendance Report") -> bytes:
        """
        Generates formatted PDF report using ReportLab.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#0F172A'),
            alignment=1,
            spaceAfter=10
        )

        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=10,
            textColor=colors.HexColor('#64748B'),
            alignment=1,
            spaceAfter=20
        )

        elements = []
        elements.append(Paragraph(f"Official College Register — {title}", title_style))
        elements.append(Paragraph(f"Generated on {datetime.now().strftime('%B %d, %Y %I:%M %p')}", sub_style))

        # Table Headers
        table_data = [["S.No", "Roll Number", "Name", "Dept", "Sec", "Subject", "Status", "Date"]]
        
        for i, item in enumerate(data[:300], 1): # Cap at 300 rows for single PDF stream
            table_data.append([
                str(i),
                str(item.get("roll_number", "")),
                str(item.get("student_name", ""))[:18],
                str(item.get("department", "")),
                str(item.get("section", "")),
                str(item.get("subject", ""))[:15],
                str(item.get("status", "")),
                str(item.get("date", ""))
            ])

        t = Table(table_data, colWidths=[35, 75, 110, 45, 45, 90, 55, 70])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8FAFC')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))

        elements.append(t)
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
