import os
import re
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Optional, Tuple, Dict, Any, List
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.models import (
    TeacherAssignment, Teacher, Subject, Section, Department, 
    AcademicYear, Student, AttendanceSession, AttendanceRecord
)
from app.services.excel_service import ExcelAttendanceService

logger = logging.getLogger("snist_erp.register_service")

def _sanitize_name(val: str) -> str:
    """Sanitizes strings for safe filenames."""
    if not val:
        return "UNKNOWN"
    return re.sub(r'[^a-zA-Z0-9_\-]', '_', val.strip()).strip('_')

def get_register_directory() -> str:
    """Ensures and returns the directory where class attendance registers reside."""
    reg_dir = os.path.join(settings.BACKEND_DIR, "data", "registers")
    os.makedirs(reg_dir, exist_ok=True)
    return reg_dir

def get_assignment_register_info(assignment: TeacherAssignment) -> Dict[str, Any]:
    """Returns metadata about the assignment's dedicated Excel register."""
    reg_dir = get_register_directory()
    file_path = assignment.excel_file_path
    file_name = assignment.excel_file_name

    exists = bool(file_path and os.path.exists(file_path))
    if not file_name:
        t_code = _sanitize_name(assignment.teacher.teacher_code if assignment.teacher else "FAC")
        sec_name = _sanitize_name(assignment.section.name if assignment.section else "SEC")
        sub_code = _sanitize_name(assignment.subject.code if assignment.subject else "SUB")
        file_name = f"Register_{assignment.id}_{t_code}_{sec_name}_{sub_code}.xlsx"

    if not file_path:
        file_path = os.path.join(reg_dir, file_name)

    return {
        "assignment_id": assignment.id,
        "file_name": file_name,
        "file_path": file_path,
        "exists": exists or os.path.exists(file_path)
    }

def generate_class_attendance_register(db: Session, assignment_id: int, overwrite: bool = False) -> str:
    """
    Generates an official SNIST attendance register Excel spreadsheet for a specific class assignment.
    Pre-populates:
      - College header & Department header
      - Class metadata (Academic Year, Year/Sem, Section, Subject, Faculty)
      - All enrolled students in this Section (S.No, Roll Number, Student Name, Agency)
      - Existing session dates and marked attendance if sessions already took place
    Saves to backend/data/registers/ and updates TeacherAssignment.
    """
    assignment = db.query(TeacherAssignment).filter(TeacherAssignment.id == assignment_id).first()
    if not assignment:
        raise ValueError(f"TeacherAssignment with id {assignment_id} not found.")

    teacher = assignment.teacher
    subject = assignment.subject
    section = assignment.section

    reg_dir = get_register_directory()
    t_code = _sanitize_name(teacher.teacher_code if teacher else "FAC")
    sec_name = _sanitize_name(section.name if section else "SEC")
    sub_code = _sanitize_name(subject.code if subject else "SUB")
    file_name = f"Register_{assignment.id}_{t_code}_{sec_name}_{sub_code}.xlsx"
    target_path = os.path.join(reg_dir, file_name)

    if os.path.exists(target_path) and not overwrite:
        # Existing register found and overwrite not forced
        if assignment.excel_file_path != target_path or assignment.excel_file_name != file_name:
            assignment.excel_file_path = target_path
            assignment.excel_file_name = file_name
            db.commit()
        return target_path

    # Query all students enrolled in this class's section
    students = db.query(Student).filter(Student.section_id == assignment.section_id).order_by(Student.roll_number.asc()).all()

    dept_name = section.department.name if (section and section.department) else "COMPUTER SCIENCE & ENGINEERING"
    year_name = section.academic_year.name if (section and section.academic_year) else "3rd Year"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Register"

    # Styling Palettes (SNIST Official Navy & Slate)
    navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    date_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    sub_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")

    white_title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    white_sub_font = Font(name="Calibri", size=11, bold=True, color="E2E8F0")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, bold=False, color="000000")
    bold_data_font = Font(name="Calibri", size=10, bold=True, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Row 1: College Header
    ws.merge_cells("A1:K1")
    ws["A1"] = "SREENIDHI INSTITUTE OF SCIENCE AND TECHNOLOGY"
    ws["A1"].font = white_title_font
    ws["A1"].fill = navy_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Row 2: Department Header
    ws.merge_cells("A2:K2")
    ws["A2"] = f"DEPARTMENT OF {dept_name.upper()} — OFFICIAL CLASS ATTENDANCE REGISTER"
    ws["A2"].font = white_sub_font
    ws["A2"].fill = navy_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Row 3 & 4: Class & Faculty Metadata
    ws["A3"] = "Academic Year:"
    ws["B3"] = year_name
    ws["D3"] = "Section:"
    ws["E3"] = section.name if section else "-"
    ws["G3"] = "Dept:"
    ws["H3"] = dept_name

    ws["A4"] = "Subject:"
    ws["B4"] = f"{subject.name} ({subject.code})" if subject else "-"
    ws["D4"] = "Faculty:"
    ws["E4"] = teacher.name if teacher else "-"
    ws["G4"] = "Faculty Code:"
    ws["H4"] = teacher.teacher_code if teacher else "-"

    for r in range(3, 5):
        ws.row_dimensions[r].height = 18
        for c in range(1, 12):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Calibri", size=10, bold=(c in [1, 4, 7]), color="000000")

    # Row 5: Date Headers (Col E onwards)
    ws.row_dimensions[5].height = 22
    for c in range(1, 5):
        cell = ws.cell(row=5, column=c)
        cell.fill = header_fill
        cell.border = thin_border
    ws.cell(row=5, column=1).value = ""

    # Row 6: Main Column Headers
    ws.row_dimensions[6].height = 24
    base_headers = [
        (1, "S.No", header_fill),
        (2, "Roll Number", header_fill),
        (3, "Student Name", header_fill),
        (4, "Agency", header_fill),
    ]
    for col_idx, h_text, fill_color in base_headers:
        cell = ws.cell(row=6, column=col_idx)
        cell.value = h_text
        cell.font = header_font
        cell.fill = fill_color
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Student Rows (Rows 7 onwards)
    start_row = 7
    for idx, s in enumerate(students, 1):
        r_idx = start_row + idx - 1
        ws.row_dimensions[r_idx].height = 19
        row_fill = PatternFill(start_color="F8FAFC" if r_idx % 2 == 0 else "FFFFFF", fill_type="solid")

        row_data = [
            (1, idx, "center", False),
            (2, s.roll_number, "center", True),
            (3, s.name, "left", False),
            (4, s.agency or "Regular", "center", False),
        ]
        for col_idx, val, align, is_bold in row_data:
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.value = val
            cell.font = bold_data_font if is_bold else data_font
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal=align, vertical="center")

    # Column Widths
    col_widths = {1: 7, 2: 18, 3: 32, 4: 14}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Save initial workbook structure
    wb.save(target_path)
    logger.info(f"Generated clean official class attendance register at {target_path}")

    # Backfill any existing sessions for this assignment into the newly created register
    try:
        sessions = db.query(AttendanceSession).filter(
            AttendanceSession.teacher_id == assignment.teacher_id,
            AttendanceSession.subject_id == assignment.subject_id,
            AttendanceSession.section_id == assignment.section_id
        ).order_by(AttendanceSession.id.asc()).all()

        for sess in sessions:
            records = db.query(AttendanceRecord).filter(AttendanceRecord.session_id == sess.id).all()
            if records:
                from datetime import datetime as dt
                try:
                    d_obj = dt.strptime(sess.session_date, "%Y-%m-%d")
                    d_formatted = d_obj.strftime("%d/%m/%Y")
                except Exception:
                    d_formatted = sess.session_date

                # Determine period count
                import re
                p_cnt = 1
                m = re.search(r'\((\d+)\s*periods?\)', sess.period or "", re.I)
                if m:
                    p_cnt = int(m.group(1))

                for r in records:
                    s_code = str(p_cnt) if r.status.value in ["PRESENT", "4"] else "A"
                    roll_str = r.student.roll_number if r.student else ""
                    if roll_str:
                        ExcelAttendanceService.record_attendance_in_excel(
                            file_path=target_path,
                            roll_number=roll_str,
                            date_str=d_formatted,
                            status_code=s_code,
                            overwrite=True
                        )
    except Exception as bf_err:
        logger.warning(f"Notice while backfilling sessions to class register {target_path}: {bf_err}")

    # Update assignment record
    assignment.excel_file_path = target_path
    assignment.excel_file_name = file_name
    db.commit()

    return target_path

def get_or_create_assignment_register(db: Session, assignment: TeacherAssignment) -> str:
    """Returns the class register path, auto-generating it if it does not yet exist."""
    if assignment.excel_file_path and os.path.exists(assignment.excel_file_path):
        return assignment.excel_file_path
    return generate_class_attendance_register(db, assignment.id)

def sync_session_to_class_register(
    db: Session,
    session_id: int,
    date_str: str,
    present_rolls: set,
    all_rolls: list,
    session_p_count: int = 1
) -> Optional[str]:
    """
    Syncs attendance marks for a completed session into its dedicated class Excel register.
    Finds the TeacherAssignment matching (teacher_id, subject_id, section_id).
    """
    session = db.query(AttendanceSession).filter(AttendanceSession.id == session_id).first()
    if not session:
        return None

    assignment = db.query(TeacherAssignment).filter(
        TeacherAssignment.teacher_id == session.teacher_id,
        TeacherAssignment.subject_id == session.subject_id,
        TeacherAssignment.section_id == session.section_id
    ).first()

    if not assignment:
        logger.info(f"No specific TeacherAssignment for session {session_id} (T={session.teacher_id}, Sub={session.subject_id}, Sec={session.section_id}).")
        return None

    try:
        class_reg_path = get_or_create_assignment_register(db, assignment)
        if not os.path.exists(class_reg_path):
            return None

        from datetime import datetime as dt
        try:
            d_obj = dt.strptime(date_str, "%Y-%m-%d")
            d_formatted = d_obj.strftime("%d/%m/%Y")
        except Exception:
            d_formatted = date_str

        for r_num in all_rolls:
            status_val = str(session_p_count) if r_num in present_rolls else "A"
            ExcelAttendanceService.record_attendance_in_excel(
                file_path=class_reg_path,
                roll_number=r_num,
                date_str=d_formatted,
                status_code=status_val,
                overwrite=True
            )
        logger.info(f"Successfully recorded attendance for session {session_id} into class register: {class_reg_path}")
        return class_reg_path
    except Exception as ex:
        logger.error(f"Error syncing session {session_id} to class register: {ex}", exc_info=True)
        return None
