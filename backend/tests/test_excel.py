import os
import openpyxl
import tempfile
import unittest
from app.services.excel_service import ExcelAttendanceService

class TestExcelService(unittest.TestCase):
    def test_excel_attendance_update(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            file_path = os.path.join(tmp_dir, "test_register.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Attendance"

            ws.append(["S.No", "Roll Number", "Student Name", "28/07/2026"])
            ws.append([1, "21311A0501", "Aarav Sharma", 4])
            ws.append([2, "21311A0502", "Aditi Verma", "A"])
            wb.save(file_path)
            wb.close()

            res = ExcelAttendanceService.record_attendance_in_excel(
                file_path=file_path,
                roll_number="21311A0502",
                date_str="29/07/2026",
                status_code="4",
                overwrite=True
            )

            self.assertEqual(res["status"], "SUCCESS")
            self.assertEqual(res["written_value"], "4")

            wb2 = openpyxl.load_workbook(file_path)
            ws2 = wb2.active
            val = ws2.cell(row=3, column=5).value
            self.assertEqual(val, 4)
            wb2.close()

if __name__ == "__main__":
    unittest.main()
