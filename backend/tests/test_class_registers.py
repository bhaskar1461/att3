import pytest
import os
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.core.database import SessionLocal
from app.models.models import User, UserRole, Teacher, Subject, Section, TeacherAssignment, Student
from app.api.auth import create_access_token
from app.services.register_service import (
    generate_class_attendance_register,
    get_or_create_assignment_register,
    sync_session_to_class_register,
    get_assignment_register_info
)

@pytest.fixture
def client():
    return TestClient(app)

@pytest.fixture
def db():
    session = SessionLocal()
    yield session
    session.close()

@pytest.fixture
def admin_token(db):
    admin = db.query(User).filter(User.role == UserRole.SUPER_ADMIN).first()
    return create_access_token(data={"sub": admin.username})

@pytest.fixture
def teacher_token(db):
    teacher = db.query(Teacher).first()
    return create_access_token(data={"sub": teacher.user.username}), teacher

def test_class_register_generation(db):
    """Verify that a class register is generated with correct sheets, headers, and student roster."""
    assignment = db.query(TeacherAssignment).first()
    assert assignment is not None

    reg_path = generate_class_attendance_register(db, assignment.id, overwrite=True)
    assert os.path.exists(reg_path)
    assert reg_path.endswith(".xlsx")

    # Inspect openpyxl workbook
    wb = openpyxl.load_workbook(reg_path, data_only=True)
    assert "Attendance Register" in wb.sheetnames
    ws = wb["Attendance Register"]

    # Verify institutional title and metadata
    assert "SREENIDHI INSTITUTE OF SCIENCE AND TECHNOLOGY" in str(ws["A1"].value)
    assert ws["A6"].value == "S.No"
    assert ws["B6"].value == "Roll Number"
    assert ws["C6"].value == "Student Name"

    # Verify roster is populated if section has students
    students_count = db.query(Student).filter(Student.section_id == assignment.section_id).count()
    if students_count > 0:
        assert ws["B7"].value is not None  # First student roll number

def test_admin_list_assignments(client, admin_token):
    """Verify admin endpoint returns rich class assignment info with excel_file_name and student_count."""
    headers = {"Authorization": f"Bearer {admin_token}"}
    res = client.get("/api/v1/admin/assignments", headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    assert len(data) > 0

    first = data[0]
    assert "id" in first
    assert "teacher_name" in first
    assert "subject_code" in first
    assert "section_name" in first
    assert "student_count" in first
    assert "excel_file_name" in first

def test_admin_download_register(client, admin_token, db):
    """Verify admin can download the class-specific register via API."""
    assignment = db.query(TeacherAssignment).first()
    headers = {"Authorization": f"Bearer {admin_token}"}
    res = client.get(f"/api/v1/admin/assignments/{assignment.id}/download-register", headers=headers)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert len(res.content) > 1000

def test_teacher_assigned_classes_includes_register_metadata(client, teacher_token):
    """Verify teacher can retrieve their distinct assigned classes with register file info."""
    token, teacher = teacher_token
    headers = {"Authorization": f"Bearer {token}"}
    res = client.get("/api/v1/teacher/assigned-classes", headers=headers)
    assert res.status_code == 200
    classes = res.json()
    assert isinstance(classes, list)
    for c in classes:
        assert "assignment_id" in c
        assert "excel_file_name" in c
        assert "has_excel_register" in c
