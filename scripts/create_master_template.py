import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure backend can be imported
backend_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "backend")
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.core.database import SessionLocal
from app.models.models import Student, Section, Department, AcademicYear

def generate_official_master_attendance_template():
    output_dir = os.path.join(backend_dir, "data", "master_templates")
    os.makedirs(output_dir, exist_ok=True)
    target_file = os.path.join(output_dir, "Official_Attendance_Register.xlsx")

    # Fetch real students for Section 1 (CS-A)
    db = SessionLocal()
    students = db.query(Student).filter(Student.section_id == 1).order_by(Student.roll_number).all()
    db.close()

    print(f"Loaded {len(students)} students from database for Section CS-A.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Register"

    # Define Palette & Fonts
    navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    date_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    light_date_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")

    white_title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    white_sub_font = Font(name="Calibri", size=11, bold=True, color="E2E8F0")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11, bold=False, color="000000")
    bold_data_font = Font(name="Calibri", size=11, bold=True, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. College Header Block
    ws.merge_cells("A1:K1")
    ws["A1"] = "SREENIDHI INSTITUTE OF SCIENCE AND TECHNOLOGY"
    ws["A1"].font = white_title_font
    ws["A1"].fill = navy_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:K2")
    ws["A2"] = "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING — OFFICIAL ATTENDANCE REGISTER"
    ws["A2"].font = white_sub_font
    ws["A2"].fill = navy_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Metadata Rows
    ws["A3"] = "Academic Year:"
    ws["B3"] = "2026-2027"
    ws["D3"] = "Year / Sem:"
    ws["E3"] = "III Year I Sem"
    ws["G3"] = "Section:"
    ws["H3"] = "CS-A"

    ws["A4"] = "Subject:"
    ws["B4"] = "Career Enhancement Training (CET)"
    ws["D4"] = "Faculty:"
    ws["E4"] = "Mrs. N. Sowjanya"
    ws["G4"] = "Period:"
    ws["H4"] = "4 Periods"

    for r in range(3, 5):
        for c in range(1, 12):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Calibri", size=10, bold=(c in [1, 4, 7]), color="000000")

    # Row 5: Date Row (Date columns from Col E onwards)
    ws.row_dimensions[5].height = 24
    for c in range(1, 5):
        cell = ws.cell(row=5, column=c)
        cell.fill = header_fill
        cell.border = thin_border
    ws.cell(row=5, column=1).value = ""
    
    # Tomorrow's Date in Column E (Column 5)
    tomorrow_date_col = 5
    tomorrow_date_header = "07/09/2026"
    date_cell = ws.cell(row=5, column=tomorrow_date_col)
    date_cell.value = tomorrow_date_header
    date_cell.font = header_font
    date_cell.fill = date_fill
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    date_cell.border = thin_border

    # Row 6: Main Column Headers
    ws.row_dimensions[6].height = 26
    headers = [
        (1, "S.No", header_fill),
        (2, "Roll Number", header_fill),
        (3, "Student Name", header_fill),
        (4, "Agency", header_fill),
        (5, 4, light_date_fill)  # Period count under date
    ]

    for col_idx, h_text, fill_color in headers:
        cell = ws.cell(row=6, column=col_idx)
        cell.value = h_text
        cell.font = header_font
        cell.fill = fill_color
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Student Rows (Rows 7 onwards)
    start_row = 7
    for idx, s in enumerate(students, 1):
        row_idx = start_row + idx - 1
        ws.row_dimensions[row_idx].height = 20
        fill = PatternFill(start_color="F8FAFC" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")

        row_data = [
            (1, idx, "center", False),
            (2, s.roll_number, "center", True),
            (3, s.name, "left", False),
            (4, s.agency or "Regular", "center", False),
            (5, None, "center", False)  # Empty ready for tomorrow's attendance
        ]

        for col_idx, val, align, is_bold in row_data:
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.font = bold_data_font if is_bold else data_font
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal=align, vertical="center")

    # Set Column Widths
    col_widths = {1: 8, 2: 18, 3: 32, 4: 14, 5: 14}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(target_file)
    print(f"Generated official master attendance register template with {len(students)} real CS-A students at: {target_file}")
    return target_file

if __name__ == "__main__":
    generate_official_master_attendance_template()
