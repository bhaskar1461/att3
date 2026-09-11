import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND_DIR = os.path.join(BASE_DIR, "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.core.database import SessionLocal
from app.models.models import User, Student, AttendanceRecord, AttendanceSession, DeviceAccountBinding, QRToken
from app.services.gsheets_service import GoogleSheetsService
from app.core.config import settings

def check():
    db = SessionLocal()
    try:
        print("=== CHECKING DB FOR BHASKAR SHARMA (23311A05Y6) ===")
        student = db.query(Student).filter(Student.roll_number == '23311A05Y6').first()
        if student:
            print(f"Student found: ID={student.id}, Name={student.name}, Section ID={student.section_id}, Roll={student.roll_number}")
            records = db.query(AttendanceRecord).filter(AttendanceRecord.student_id == student.id).all()
            print(f"Attendance records count: {len(records)}")
            for r in records[:5]:
                print(f"  Record: session_id={r.session_id}, date={r.session_date}, status={r.status}, periods={r.period_count}")
        else:
            print("No Student with roll 23311A05Y6 found.")

        user = db.query(User).filter(User.username == '23311A05Y6').first()
        if user:
            print(f"User found: ID={user.id}, Username={user.username}, Role={user.role}")
        else:
            print("No User with username 23311A05Y6 found.")

        bindings = db.query(DeviceAccountBinding).filter(DeviceAccountBinding.roll_number == '23311A05Y6').all()
        print(f"DeviceAccountBinding count: {len(bindings)}")

        # Check section 1 total students
        sec1_students = db.query(Student).filter(Student.section_id == 1).order_by(Student.roll_number).all()
        print(f"Section 1 total students: {len(sec1_students)}")
        sec1_rolls = [s.roll_number for s in sec1_students]
        if '23311A05Y6' in sec1_rolls:
            print("WARNING: 23311A05Y6 IS in Section 1!")
        else:
            print("CONFIRMED: 23311A05Y6 is NOT in Section 1.")

        # Check GSheet
        print("\n=== CHECKING GOOGLE SHEET 'CSE-CS' ===")
        import time
        creds = settings.GOOGLE_CREDENTIALS_FILE or os.path.join(BACKEND_DIR, "credentials.json")
        sheet_id = "1CDeeivsdptGtgJpgK6XN9o6AKqy6HAIucv45D0Os6ZY"
        client = GoogleSheetsService._get_client(creds)
        vals = []
        for attempt in range(3):
            try:
                sh = client.open_by_key(sheet_id)
                ws = None
                for title in ["CSE-CS", "Attendance Register"]:
                    try:
                        ws = sh.worksheet(title)
                        break
                    except Exception:
                        pass
                if not ws:
                    ws = sh.sheet1
                vals = ws.get_all_values()
                print(f"Total rows in GSheet '{ws.title}': {len(vals)}")
                break
            except Exception as e:
                print(f"GSheet attempt {attempt+1} failed: {e}")
                time.sleep(15)

        found_in_sheet = []
        for idx, row in enumerate(vals, start=1):
            row_text = " ".join(row).upper()
            if "23311A05Y6" in row_text or "BHASKAR" in row_text:
                found_in_sheet.append((idx, row[:5]))
        
        if found_in_sheet:
            print("FOUND IN GSHEET:")
            for item in found_in_sheet:
                print(f"  Row {item[0]}: {item[1]}")
        else:
            print("CONFIRMED: 23311A05Y6 / Bhaskar is NOT in Google Sheet.")

        # Check Excel
        print("\n=== CHECKING MASTER EXCEL REGISTER ===")
        excel_path = os.path.join(BACKEND_DIR, "data", "master_templates", "Official_Attendance_Register.xlsx")
        if os.path.exists(excel_path):
            import openpyxl
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws_x = wb.active
            found_in_excel = []
            for r in range(1, ws_x.max_row + 1):
                row_vals = [str(ws_x.cell(row=r, column=c).value or "") for c in range(1, 10)]
                row_text = " ".join(row_vals).upper()
                if "23311A05Y6" in row_text or "BHASKAR" in row_text:
                    found_in_excel.append((r, row_vals[:4]))
            if found_in_excel:
                print("FOUND IN EXCEL:")
                for item in found_in_excel:
                    print(f"  Row {item[0]}: {item[1]}")
            else:
                print("CONFIRMED: 23311A05Y6 / Bhaskar is NOT in Master Excel Register.")
        else:
            print(f"Excel file does not exist: {excel_path}")

    finally:
        db.close()

if __name__ == "__main__":
    check()
