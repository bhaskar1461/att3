import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Add backend directory to sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND_DIR = os.path.join(BASE_DIR, "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.core.config import settings
from app.core.database import SessionLocal
from app.models.models import Student, SystemSettings
from app.services.gsheets_service import GoogleSheetsService

def fetch_students_from_db():
    db = SessionLocal()
    try:
        students = db.query(Student).all()
        result = []
        for s in students:
            result.append({
                "id": s.id,
                "roll_number": s.roll_number,
                "name": s.name,
                "department": s.department.name if s.department else "N/A",
                "academic_year": s.academic_year.name if s.academic_year else "N/A",
                "section": s.section.name if s.section else "N/A",
                "email": s.email or "",
                "mobile": s.mobile or "",
                "agency": s.agency or "Regular"
            })
        return result
    except Exception as e:
        print(f"[!] Error fetching students from database: {e}")
        return []
    finally:
        db.close()

def generate_local_excel(students, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Roster"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    navy_fill = PatternFill(start_color="15347E", end_color="15347E", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = ["S.No", "Roll Number", "Student Name", "Department", "Academic Year", "Section", "Email", "Mobile", "Category"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, s in enumerate(students, 1):
        row = [
            idx,
            s["roll_number"],
            s["name"],
            s["department"],
            s["academic_year"],
            s["section"],
            s["email"],
            s["mobile"],
            s["agency"]
        ]
        ws.append(row)
        current_row = ws.max_row
        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = cell_font
            cell.border = thin_border

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"[+] Saved local Excel import sheet to: {output_path}")

def update_system_settings_spreadsheet_id(spreadsheet_id):
    db = SessionLocal()
    try:
        setting = db.query(SystemSettings).filter(SystemSettings.key == 'GOOGLE_SPREADSHEET_ID').first()
        if setting:
            setting.value = spreadsheet_id
        else:
            setting = SystemSettings(key='GOOGLE_SPREADSHEET_ID', value=spreadsheet_id, description='Google Sheets Spreadsheet ID')
            db.add(setting)
        db.commit()
        print(f"[+] System setting GOOGLE_SPREADSHEET_ID updated to: {spreadsheet_id}")
    except Exception as e:
        db.rollback()
        print(f"[!] Error updating system setting: {e}")
    finally:
        db.close()

def main():
    print("==================================================")
    print("   EXPORTS ALL STUDENTS TO GOOGLE SHEETS / EXCEL   ")
    print("==================================================")

    students = fetch_students_from_db()
    print(f"[*] Found {len(students)} students in the database.")

    if not students:
        print("[!] No students found to export.")
        return

    # Always generate local formatted Excel file
    excel_path = os.path.join(settings.OUTPUT_EXCEL_DIR, "Students_Google_Sheet_Import.xlsx")
    generate_local_excel(students, excel_path)

    # Check for Google Service Account credentials
    creds_file = settings.GOOGLE_CREDENTIALS_FILE or os.getenv("GOOGLE_CREDENTIALS_FILE")
    
    # Check if a credentials.json file exists in backend/ directory as fallback
    if not creds_file:
        fallback_json = os.path.join(BACKEND_DIR, "credentials.json")
        if os.path.exists(fallback_json):
            creds_file = fallback_json

    if creds_file and os.path.exists(creds_file):
        print(f"[*] Found Google Credentials file at: {creds_file}")
        print("[*] Creating Google Sheet and uploading student records...")

        result = GoogleSheetsService.create_and_populate_student_sheet(
            credentials_json=creds_file,
            students=students,
            title="Student Roster - AI QR Attendance System",
            share_email="anyone" # Shared publicly via view link if supported
        )

        if result.get("success"):
            sp_id = result["spreadsheet_id"]
            url = result["url"]
            print("\n" + "="*60)
            print("[SUCCESS] GOOGLE SHEET CREATED SUCCESSFULLY!")
            print(f"Spreadsheet ID : {sp_id}")
            print(f"Spreadsheet URL: {url}")
            print("="*60 + "\n")

            update_system_settings_spreadsheet_id(sp_id)
        else:
            print(f"[!] Failed to create Google Sheet: {result.get('error')}")
    else:
        print("\n" + "-"*60)
        print("[NOTE] Google Credentials file not provided yet.")
        print("To create/sync live on Google Cloud:")
        print(" 1. Place your Google Service Account key file at 'backend/credentials.json'")
        print("    or set GOOGLE_CREDENTIALS_FILE in .env")
        print(" 2. Re-run this script.")
        print("")
        print(f"Local file created ready for 1-click upload to Google Drive:")
        print(f" -> {excel_path}")
        print("-"*60 + "\n")

if __name__ == "__main__":
    main()
