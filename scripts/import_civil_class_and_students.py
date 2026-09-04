import os
import sys
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND_DIR = os.path.join(BASE_DIR, "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.core.database import SessionLocal
from app.models.models import (
    Department, AcademicYear, Section, Subject, Teacher, 
    TeacherAssignment, Student, User, UserRole
)
from app.core.security import get_password_hash

EXCEL_FILE = os.path.join(BASE_DIR, "Civil  III- I - Sample Attendance Sheet.xlsx")

def main():
    print(f"Loading Excel file: {EXCEL_FILE}")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb.active

    db = SessionLocal()
    try:
        # 1. Ensure Civil Department exists
        civil_dept = db.query(Department).filter(Department.code == "CIVIL").first()
        if not civil_dept:
            civil_dept = Department(code="CIVIL", name="Civil Engineering")
            db.add(civil_dept)
            db.flush()
            print(f"Created Department: CIVIL (ID: {civil_dept.id})")
        else:
            print(f"Found Department: CIVIL (ID: {civil_dept.id})")

        # 2. Ensure 3rd Year exists
        yr_3 = db.query(AcademicYear).filter(AcademicYear.name == "3rd Year").first()
        if not yr_3:
            yr_3 = AcademicYear(name="3rd Year")
            db.add(yr_3)
            db.flush()
            print(f"Created AcademicYear: 3rd Year (ID: {yr_3.id})")
        else:
            print(f"Found AcademicYear: 3rd Year (ID: {yr_3.id})")

        # 3. Ensure Section Civil III-I exists
        civil_sec = db.query(Section).filter(Section.name == "Civil III-I").first()
        if not civil_sec:
            civil_sec = Section(
                name="Civil III-I",
                department_id=civil_dept.id,
                academic_year_id=yr_3.id
            )
            db.add(civil_sec)
            db.flush()
            print(f"Created Section: Civil III-I (ID: {civil_sec.id})")
        else:
            print(f"Found Section: Civil III-I (ID: {civil_sec.id})")

        # 4. Ensure Subject CET exists
        cet_subj = db.query(Subject).filter(Subject.code == "CE301").first()
        if not cet_subj:
            cet_subj = Subject(
                code="CE301",
                name="Career Enhancement Training (CET) - Civil",
                department_id=civil_dept.id,
                academic_year_id=yr_3.id
            )
            db.add(cet_subj)
            db.flush()
            print(f"Created Subject: CE301 (ID: {cet_subj.id})")
        else:
            print(f"Found Subject: CE301 (ID: {cet_subj.id})")

        # 5. Assign Teacher (Prof. Srinivas Rao) to Civil III-I
        teacher = db.query(Teacher).filter(Teacher.name.like("%Srinivas%")).first()
        if not teacher:
            teacher = db.query(Teacher).first()

        if teacher:
            assignment = db.query(TeacherAssignment).filter(
                TeacherAssignment.teacher_id == teacher.id,
                TeacherAssignment.subject_id == cet_subj.id,
                TeacherAssignment.section_id == civil_sec.id
            ).first()
            if not assignment:
                assignment = TeacherAssignment(
                    teacher_id=teacher.id,
                    subject_id=cet_subj.id,
                    section_id=civil_sec.id
                )
                db.add(assignment)
                db.flush()
                print(f"Assigned Teacher '{teacher.name}' to 'Civil III-I - {cet_subj.name}'")
            else:
                print(f"Teacher '{teacher.name}' already assigned to 'Civil III-I - {cet_subj.name}'")

        # 6. Parse and Import Students from Excel
        imported_count = 0
        updated_count = 0
        max_row = ws.max_row

        for r in range(7, max_row + 1):
            roll_val = ws.cell(r, 2).value
            name_val = ws.cell(r, 3).value
            agency_val = ws.cell(r, 4).value or "CIVIL(CET)"

            if not roll_val:
                continue

            roll = str(roll_val).strip().upper()
            name = str(name_val).strip() if name_val else f"Student {roll}"

            # Check if student exists
            existing_student = db.query(Student).filter(Student.roll_number == roll).first()
            if existing_student:
                # Update section to Civil III-I
                existing_student.section_id = civil_sec.id
                existing_student.department_id = civil_dept.id
                existing_student.name = name
                existing_student.agency = str(agency_val).strip()
                updated_count += 1
            else:
                # Create user account
                user = db.query(User).filter(User.username == roll).first()
                if not user:
                    user = User(
                        username=roll,
                        email=f"{roll.lower()}@snist.edu.in",
                        password_hash=get_password_hash("student123"),
                        role=UserRole.STUDENT
                    )
                    db.add(user)
                    db.flush()

                student = Student(
                    user_id=user.id,
                    roll_number=roll,
                    name=name,
                    department_id=civil_dept.id,
                    academic_year_id=yr_3.id,
                    section_id=civil_sec.id,
                    email=f"{roll.lower()}@snist.edu.in",
                    agency=str(agency_val).strip()
                )
                db.add(student)
                imported_count += 1

        db.commit()
        print(f"SUCCESS! Imported {imported_count} new Civil students, updated {updated_count} existing students into 'Civil III-I'.")
        
        # Verify total students in Civil III-I
        total_civil = db.query(Student).filter(Student.section_id == civil_sec.id).count()
        print(f"Total students currently in Civil III-I section: {total_civil}")

    except Exception as ex:
        db.rollback()
        print(f"Error during import: {ex}")
        raise
    finally:
        db.close()

if __name__ == "__main__":
    main()
