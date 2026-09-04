import os
import sys
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND_DIR = os.path.join(BASE_DIR, "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.core.database import SessionLocal, engine, Base
from app.models.models import (
    User, UserRole, Department, AcademicYear, Section, Subject,
    Teacher, TeacherAssignment, Student, AttendanceRecord,
    DeviceAccountBinding, QRToken
)
from app.core.security import get_password_hash

EXCEL_PATH = os.path.join(BASE_DIR, "CSE-CS  III - I.xlsx")

def parse_students_from_excel(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = wb["CSE-CS"]
    students = []
    
    for r in range(7, sheet.max_row + 1):
        sno = sheet.cell(row=r, column=1).value
        roll = sheet.cell(row=r, column=2).value
        name = sheet.cell(row=r, column=3).value
        gender = sheet.cell(row=r, column=4).value
        sec = sheet.cell(row=r, column=5).value
        agency = sheet.cell(row=r, column=6).value or "CS(CET)"

        if roll and str(roll).strip():
            clean_roll = str(roll).strip().upper()
            clean_name = str(name).strip() if name else f"Student {clean_roll}"
            clean_sec = str(sec).strip() if sec else "A"
            clean_gender = str(gender).strip() if gender else ""
            clean_agency = str(agency).strip()

            students.append({
                "sno": sno,
                "roll": clean_roll,
                "name": clean_name,
                "gender": clean_gender,
                "section": clean_sec,
                "agency": clean_agency
            })
    return students

def execute_reseed():
    print(f"Loading Excel student data from: {EXCEL_PATH}")
    parsed_students = parse_students_from_excel(EXCEL_PATH)
    print(f"Found {len(parsed_students)} students in Excel file.")

    db = SessionLocal()
    try:
        print("\n--- PHASE 1: Removing Old Student Data ---")
        # 1. Gather all existing student IDs and usernames
        old_students = db.query(Student).all()
        old_student_ids = [s.id for s in old_students]
        old_student_user_ids = [s.user_id for s in old_students if s.user_id]
        old_student_rolls = [s.roll_number for s in old_students]

        print(f"Identified {len(old_students)} old student records to delete.")

        # Delete dependent attendance records
        if old_student_ids:
            deleted_records = db.query(AttendanceRecord).filter(AttendanceRecord.student_id.in_(old_student_ids)).delete(synchronize_session=False)
            print(f"Deleted {deleted_records} old attendance records.")

        # Delete dependent QR tokens
        if old_student_ids:
            deleted_tokens = db.query(QRToken).filter(QRToken.student_id.in_(old_student_ids)).delete(synchronize_session=False)
            print(f"Deleted {deleted_tokens} old QR tokens.")

        # Delete device account bindings for old roll numbers
        if old_student_rolls:
            deleted_bindings = db.query(DeviceAccountBinding).filter(DeviceAccountBinding.roll_number.in_(old_student_rolls)).delete(synchronize_session=False)
            print(f"Deleted {deleted_bindings} old device bindings.")

        # Delete student profiles
        deleted_students = db.query(Student).delete(synchronize_session=False)
        print(f"Deleted {deleted_students} student profiles.")

        # Delete student user accounts (STRICTLY preserve SUPER_ADMIN and TEACHER)
        deleted_users = db.query(User).filter(User.role == UserRole.STUDENT).delete(synchronize_session=False)
        print(f"Deleted {deleted_users} student user accounts.")
        db.flush()

        print("\n--- PHASE 2: Configuring Department, Academic Year, and Section ---")
        # Department: CSE-CS
        dept = db.query(Department).filter(Department.code == "CSE-CS").first()
        if not dept:
            dept = Department(code="CSE-CS", name="Computer Science and Engineering (Cyber Security)")
            db.add(dept)
            db.flush()
            print(f"Created Department: {dept.code} - {dept.name} (ID: {dept.id})")
        else:
            print(f"Using existing Department: {dept.code} (ID: {dept.id})")

        # Academic Year: 3rd Year
        yr_3 = db.query(AcademicYear).filter(AcademicYear.name == "3rd Year").first()
        if not yr_3:
            yr_3 = AcademicYear(name="3rd Year")
            db.add(yr_3)
            db.flush()
            print(f"Created Academic Year: 3rd Year (ID: {yr_3.id})")
        else:
            print(f"Using Academic Year: 3rd Year (ID: {yr_3.id})")

        # Section: CSE-CS III-I
        sec = db.query(Section).filter(Section.name == "CSE-CS III-I").first()
        if not sec:
            sec = Section(
                name="CSE-CS III-I",
                department_id=dept.id,
                academic_year_id=yr_3.id
            )
            db.add(sec)
            db.flush()
            print(f"Created Section: CSE-CS III-I (ID: {sec.id})")
        else:
            sec.department_id = dept.id
            sec.academic_year_id = yr_3.id
            db.flush()
            print(f"Using existing Section: CSE-CS III-I (ID: {sec.id})")

        # Subject: Career Enhancement Training (CET) - CSE-CS
        subject = db.query(Subject).filter(Subject.code == "CET-CS").first()
        if not subject:
            subject = Subject(
                code="CET-CS",
                name="Career Enhancement Training (CET) - CSE-CS",
                department_id=dept.id,
                academic_year_id=yr_3.id
            )
            db.add(subject)
            db.flush()
            print(f"Created Subject: CET-CS (ID: {subject.id})")
        else:
            print(f"Using Subject: CET-CS (ID: {subject.id})")

        # Assign Teachers to this class so they can take attendance
        teachers = db.query(Teacher).all()
        for t in teachers:
            assignment = db.query(TeacherAssignment).filter(
                TeacherAssignment.teacher_id == t.id,
                TeacherAssignment.subject_id == subject.id,
                TeacherAssignment.section_id == sec.id
            ).first()
            if not assignment:
                assignment = TeacherAssignment(
                    teacher_id=t.id,
                    subject_id=subject.id,
                    section_id=sec.id
                )
                db.add(assignment)
                db.flush()
                print(f"Assigned Teacher '{t.name}' to 'CSE-CS III-I - {subject.name}'")

        print("\n--- PHASE 3: Seeding 50 CSE-CS Students & Users ---")
        student_pw_hash = get_password_hash("student123")
        created_count = 0

        for s_data in parsed_students:
            roll = s_data["roll"]
            name = s_data["name"]
            agency = s_data["agency"]

            user = User(
                username=roll,
                email=f"{roll.lower()}@snist.edu.in",
                password_hash=student_pw_hash,
                role=UserRole.STUDENT
            )
            db.add(user)
            db.flush()

            student = Student(
                user_id=user.id,
                roll_number=roll,
                name=name,
                department_id=dept.id,
                academic_year_id=yr_3.id,
                section_id=sec.id,
                email=f"{roll.lower()}@snist.edu.in",
                agency=agency
            )
            db.add(student)
            created_count += 1

        db.commit()
        print(f"\n✅ SUCCESS: Successfully seeded {created_count} students into Section 'CSE-CS III-I'!")
        
        # Verification counts
        total_students = db.query(Student).count()
        total_student_users = db.query(User).filter(User.role == UserRole.STUDENT).count()
        print(f"Total students now in database: {total_students}")
        print(f"Total student user accounts: {total_student_users}")
        print(f"Total staff accounts preserved: {db.query(User).filter(User.role != UserRole.STUDENT).count()}")

    except Exception as e:
        db.rollback()
        print(f"❌ Error during reseed: {e}")
        raise
    finally:
        db.close()

if __name__ == "__main__":
    execute_reseed()
