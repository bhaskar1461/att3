import os
import sys
import pickle
import openpyxl
import datetime
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
import gspread

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

TOKEN_PATH = os.path.join(BASE_DIR, "token.pickle")
EXCEL_PATH = os.path.join(BASE_DIR, "Civil  III- I - Sample Attendance Sheet.xlsx")

def main():
    print(f"Loading credentials from: {TOKEN_PATH}")
    with open(TOKEN_PATH, "rb") as f:
        creds = pickle.load(f)

    if not creds.valid or creds.expired:
        print("Refreshing expired credentials...")
        creds.refresh(Request())
        with open(TOKEN_PATH, "wb") as f:
            pickle.dump(creds, f)
        backend_token = os.path.join(BASE_DIR, "backend", "token.pickle")
        if os.path.exists(os.path.dirname(backend_token)):
            with open(backend_token, "wb") as f:
                pickle.dump(creds, f)
        print("Credentials refreshed successfully.")

    client = gspread.authorize(creds)
    drive_service = build("drive", "v3", credentials=creds)

    print(f"Reading Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active

    max_row = ws.max_row
    max_col = 19  # Columns A to S

    grid = []
    for r in range(1, max_row + 1):
        row_data = []
        for c in range(1, max_col + 1):
            val = ws.cell(r, c).value
            if isinstance(val, (datetime.datetime, datetime.date)):
                val = f"{val.day}/{val.month}/{str(val.year)[-2:]}"
            elif val is None:
                val = ""
            else:
                val = str(val).strip()
                if val == "4":
                    val = 4
                elif val.isdigit() and c == 1:
                    val = int(val)
            row_data.append(val)
        grid.append(row_data)

    sheet_title = "Civil III-I - Attendance Register"
    print(f"Creating new Google Spreadsheet: '{sheet_title}'...")
    spreadsheet = client.create(sheet_title)
    spreadsheet_id = spreadsheet.id
    spreadsheet_url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit"
    print(f"Spreadsheet Created! ID: {spreadsheet_id}")
    print(f"URL: {spreadsheet_url}")

    # Make spreadsheet editable by anyone with link
    try:
        drive_service.permissions().create(
            fileId=spreadsheet_id,
            body={"type": "anyone", "role": "writer"},
            fields="id"
        ).execute()
        print("Permissions set: Anyone with link can edit.")
    except Exception as e:
        print(f"Permission warning: {e}")

    # Primary worksheet
    worksheet = spreadsheet.sheet1
    worksheet.update_title("Attendance Register")
    worksheet.resize(rows=max(max_row + 10, 100), cols=max_col + 5)

    print(f"Uploading {len(grid)} rows and {max_col} columns of data...")
    worksheet.update("A1:S" + str(len(grid)), grid, value_input_option="USER_ENTERED")

    # Apply batch formatting
    sheet_id = worksheet.id
    target_cols = max_col
    num_rows = len(grid)

    requests = [
        # 1. Merge Header Rows A1:S1, A2:S2, A3:S3, A4:S4, A5:D5
        {
            "mergeCells": {
                "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "mergeType": "MERGE_ALL"
            }
        },
        {
            "mergeCells": {
                "range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": 2, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "mergeType": "MERGE_ALL"
            }
        },
        {
            "mergeCells": {
                "range": {"sheetId": sheet_id, "startRowIndex": 2, "endRowIndex": 3, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "mergeType": "MERGE_ALL"
            }
        },
        {
            "mergeCells": {
                "range": {"sheetId": sheet_id, "startRowIndex": 3, "endRowIndex": 4, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "mergeType": "MERGE_ALL"
            }
        },
        {
            "mergeCells": {
                "range": {"sheetId": sheet_id, "startRowIndex": 4, "endRowIndex": 5, "startColumnIndex": 0, "endColumnIndex": 4},
                "mergeType": "MERGE_ALL"
            }
        },

        # 2. Column Dimensions
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 1}, "properties": {"pixelSize": 55}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": 1, "endIndex": 2}, "properties": {"pixelSize": 130}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": 2, "endIndex": 3}, "properties": {"pixelSize": 240}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": 3, "endIndex": 4}, "properties": {"pixelSize": 95}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": 4, "endIndex": target_cols}, "properties": {"pixelSize": 72}, "fields": "pixelSize"}},

        # 3. Row heights for top header block
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": 0, "endIndex": 6}, "properties": {"pixelSize": 36}, "fields": "pixelSize"}},

        # 4. Styling Row 1: Navy Blue Header
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.08, "green": 0.20, "blue": 0.49},
                        "textFormat": {"bold": True, "fontSize": 13, "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        # Row 2: Department Title
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": 2, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.12, "green": 0.26, "blue": 0.58},
                        "textFormat": {"bold": True, "fontSize": 11, "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        # Row 3: CET Title
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 2, "endRowIndex": 3, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.90, "green": 0.93, "blue": 0.98},
                        "textFormat": {"bold": True, "fontSize": 10, "foregroundColor": {"red": 0.08, "green": 0.20, "blue": 0.49}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        # Row 4: Agency info
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 3, "endRowIndex": 4, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.95, "green": 0.96, "blue": 0.98},
                        "textFormat": {"bold": True, "fontSize": 9, "foregroundColor": {"red": 0.2, "green": 0.2, "blue": 0.2}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        # Row 5: Batch info + Dates
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 4, "endRowIndex": 5, "startColumnIndex": 0, "endColumnIndex": 4},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.88, "green": 0.92, "blue": 0.96},
                        "textFormat": {"bold": True, "fontSize": 9, "foregroundColor": {"red": 0.08, "green": 0.20, "blue": 0.49}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 4, "endRowIndex": 5, "startColumnIndex": 4, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.95, "green": 0.97, "blue": 1.0},
                        "textFormat": {"bold": True, "fontSize": 9, "foregroundColor": {"red": 0.08, "green": 0.20, "blue": 0.49}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        # Row 6: Column Headers (SNO, ROLL NO, NAME, Agency, 4, 4...)
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 5, "endRowIndex": 6, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.22, "green": 0.32, "blue": 0.48},
                        "textFormat": {"bold": True, "fontSize": 9, "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        # Data Rows (Row 7 to end): General alignments
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 6, "endRowIndex": num_rows, "startColumnIndex": 0, "endColumnIndex": 1},
                "cell": {
                    "userEnteredFormat": {
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(horizontalAlignment,verticalAlignment)"
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 6, "endRowIndex": num_rows, "startColumnIndex": 1, "endColumnIndex": 2},
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {"bold": True, "fontFamily": "Courier New"},
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 6, "endRowIndex": num_rows, "startColumnIndex": 2, "endColumnIndex": 3},
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {"bold": True},
                        "horizontalAlignment": "LEFT",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(textFormat,horizontalAlignment,verticalAlignment)"
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 6, "endRowIndex": num_rows, "startColumnIndex": 3, "endColumnIndex": target_cols},
                "cell": {
                    "userEnteredFormat": {
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE"
                    }
                },
                "fields": "userEnteredFormat(horizontalAlignment,verticalAlignment)"
            }
        },

        # Borders across the whole table
        {
            "updateBorders": {
                "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": num_rows, "startColumnIndex": 0, "endColumnIndex": target_cols},
                "top": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.85, "blue": 0.9}},
                "bottom": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.85, "blue": 0.9}},
                "left": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.85, "blue": 0.9}},
                "right": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.85, "blue": 0.9}},
                "innerHorizontal": {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.88, "blue": 0.92}},
                "innerVertical": {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.88, "blue": 0.92}}
            }
        }
    ]

    print("Applying styling, column widths, and cell merges...")
    spreadsheet.batch_update({"requests": requests})
    print("\n========================================================")
    print("SUCCESS_SPREADSHEET_ID:" + spreadsheet_id)
    print("SUCCESS_SPREADSHEET_URL:" + spreadsheet_url)
    print("========================================================\n")

if __name__ == "__main__":
    main()
